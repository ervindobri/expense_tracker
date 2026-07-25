"""
Management command to import historical Entry data from the yearly
budget spreadsheet (e.g. Spent.xlsx).

Usage:
    python manage.py import_expenses /path/to/Spent.xlsx
    python manage.py import_expenses /path/to/Spent.xlsx --year 2026
    python manage.py import_expenses /path/to/Spent.xlsx --dry-run

Place this file at:
    yourapp/management/commands/import_expenses.py

(Django needs empty __init__.py files in both yourapp/management/ and
yourapp/management/commands/ for the command to be discovered.)
"""

import calendar
from datetime import datetime

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

# Replace with your actual app's import path
from tracker.models import Category, CategoryType, Entry

CATEGORY_COL = 4       # column D — category name
MONTH_TITLE_COL = 3    # column C — month name, one row above the header row
WEEK_COLS = [5, 7, 9, 11, 13]  # columns E, G, I, K, M — WEEK1..WEEK5

# First day of each week range, per the sheet's grouping:
# week1: 1-6, week2: 7-13, week3: 14-20, week4: 21-27, week5: 28-31
WEEK_START_DAYS = {1: 1, 2: 7, 3: 14, 4: 21, 5: 28}


class Command(BaseCommand):
    help = "Import Entry records from a yearly budget spreadsheet (.xlsx)."

    def add_arguments(self, parser):
        parser.add_argument("file_path", type=str, help="Path to the .xlsx file")
        parser.add_argument(
            "--year",
            type=int,
            default=timezone.now().year,
            help="Which year's sheet to import (default: current year)",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Preview what would be imported without writing to the database",
        )

    def handle(self, *args, **options):
        file_path = options["file_path"]
        year = options["year"]
        dry_run = options["dry_run"]

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {file_path}")
        except Exception as exc:
            raise CommandError(f"Could not open workbook: {exc}")

        sheet_name = str(year)
        if sheet_name not in wb.sheetnames:
            raise CommandError(
                f"No sheet named '{sheet_name}'. Available sheets: {wb.sheetnames}"
            )
        ws = wb[sheet_name]

        created_count = 0
        skipped_categories = set()
        category_cache = {}

        with transaction.atomic():
            for header_row in self._find_header_rows(ws):
                month_name = ws.cell(row=header_row - 1, column=MONTH_TITLE_COL).value
                try:
                    month_num = datetime.strptime(str(month_name).strip(), "%B").month
                except (ValueError, AttributeError):
                    self.stdout.write(
                        self.style.WARNING(
                            f"Row {header_row}: couldn't parse month name "
                            f"'{month_name}', skipping this block."
                        )
                    )
                    continue

                days_in_month = calendar.monthrange(year, month_num)[1]
                created_count += self._process_month_block(
                    ws=ws,
                    start_row=header_row + 1,
                    year=year,
                    month_num=month_num,
                    month_name=month_name,
                    days_in_month=days_in_month,
                    category_cache=category_cache,
                    skipped_categories=skipped_categories,
                    dry_run=dry_run,
                )

            if dry_run:
                # Roll back — dry run should never persist anything
                transaction.set_rollback(True)

        if skipped_categories:
            self.stdout.write(
                self.style.WARNING(
                    "New categories not previously in your DB were created: "
                    f"{sorted(skipped_categories)}"
                )
            )

        if dry_run:
            self.stdout.write(
                self.style.SUCCESS(
                    f"Dry run complete — {created_count} entries would be created."
                )
            )
        else:
            self.stdout.write(
                self.style.SUCCESS(f"Import complete — {created_count} entries created.")
            )

    def _find_header_rows(self, ws):
        """Every row where column D == 'CATEGORY' marks the start of a month block."""
        header_rows = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=CATEGORY_COL):
            if row[CATEGORY_COL - 1].value == "CATEGORY":
                header_rows.append(row[0].row)
        return header_rows

    def _process_month_block(
        self,
        ws,
        start_row,
        year,
        month_num,
        month_name,
        days_in_month,
        category_cache,
        skipped_categories,
        dry_run,
    ):
        created = 0
        mode = None  # CategoryType.EXPENSE or CategoryType.INCOME, set by section headers
        row = start_row

        while True:
            cell_value = ws.cell(row=row, column=CATEGORY_COL).value
            if cell_value is None:
                break  # blank row ends this month's block

            label = str(cell_value).strip()
            lowered = label.lower()

            if "expenses" in lowered:
                mode = CategoryType.EXPENSE
            elif "incomes" in lowered:
                mode = CategoryType.INCOME
            elif mode is not None:
                category = self._get_category(
                    category_cache, label, mode, skipped_categories, dry_run
                )

                for week_num, col in zip(range(1, 6), WEEK_COLS):
                    amount = ws.cell(row=row, column=col).value
                    if not amount:  # skips both None and 0
                        continue

                    day = min(WEEK_START_DAYS[week_num], days_in_month)
                    naive_dt = datetime(year, month_num, day)
                    added_date = (
                        timezone.make_aware(naive_dt) if settings.USE_TZ else naive_dt
                    )

                    if dry_run:
                        self.stdout.write(
                            f"  [{month_name}] {label} — week {week_num}: "
                            f"{amount} on {added_date.date()}"
                        )
                        created += 1
                    else:
                        _, was_created = Entry.objects.get_or_create(
                            category=category,
                            added_date=added_date,
                            amount=amount,
                        )
                        if was_created:
                            created += 1

            row += 1

        return created

    def _get_category(self, cache, name, category_type, skipped_categories, dry_run):
        key = (name, category_type)
        if key in cache:
            return cache[key]

        if dry_run:
            cache[key] = None
            return None

        category, was_created = Category.objects.get_or_create(
            name=name,
            category_type=category_type,
            defaults={"pub_date": timezone.now()},
        )
        if was_created:
            skipped_categories.add(name)

        cache[key] = category
        return category
